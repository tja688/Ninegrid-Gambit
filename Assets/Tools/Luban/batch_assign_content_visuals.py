#!/usr/bin/env python3
"""Batch-assign icon/face visual_ids + visual_asset legacy keys for bootstrap art pass."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent
DATAS_DIR = SCRIPT_DIR / "Datas"
PROJECT_ROOT = SCRIPT_DIR.parents[2]

CONTENT_VISUAL_PATH = DATAS_DIR / "content_visual.xlsx"
VISUAL_ASSET_PATH = DATAS_DIR / "visual_asset.xlsx"

MONSTER_ICON_ATLAS = "Assets/Arts/Images/Multiple/pixel_art_card_pack 1.png"
MONSTER_FACE_ATLAS = "Assets/Arts/Images/Multiple/pixel_art_card_pack.png"
FACE_SPRITE_NAMES = [f"pixel_art_card_pack_{i}" for i in range(1, 11)]

ITEMS_DIR = PROJECT_ROOT / "Assets/Arts/Images/Png/Items"
ICONS_DIR = PROJECT_ROOT / "Assets/Arts/Images/Png/Icons"

CV_SHEET = "TbContentVisual"
VA_SHEET = "TbVisualAsset"


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_atlas_sprites(meta_path: Path) -> list[str]:
    if not meta_path.exists():
        return []
    text = meta_path.read_text(encoding="utf-8", errors="ignore")
    names: list[str] = []
    for match in re.finditer(r"second:\s*(.+?)\s*$", text, re.MULTILINE):
        name = match.group(1).strip()
        if name and name not in names:
            names.append(name)
    return names


def collect_png_sprites(root: Path) -> dict[str, str]:
    """Map normalized stem -> legacy asset key Assets/...#spriteName."""
    mapping: dict[str, str] = {}
    if not root.exists():
        return mapping
    for png in root.rglob("*.png"):
        rel = png.relative_to(PROJECT_ROOT).as_posix()
        stem = png.stem.lower()
        sprite_name = png.stem
        key = f"{rel}#{sprite_name}"
        mapping[stem] = key
        # also register path segments for partial match
        for part in stem.replace("-", "_").split("_"):
            if len(part) >= 3 and part not in mapping:
                mapping[part] = key
    return mapping


def visual_id(slot: str, content_id: str) -> str:
    return f"visual.{slot}.{content_id}"


def legacy_key(atlas_path: str, sprite_name: str) -> str:
    return f"{atlas_path}#{sprite_name}"


def tokenize_content_id(content_id: str) -> list[str]:
    raw = content_id.split(".")[-1]
    parts = re.split(r"[_\-\s]+", raw.lower())
    return [p for p in parts if len(p) >= 2]


def pick_from_pool(content_id: str, pool: list[str]) -> str:
    if not pool:
        return ""
    index = sum(ord(c) for c in content_id) % len(pool)
    return pool[index]


def fuzzy_pick(content_id: str, sprite_map: dict[str, str], fallback_pool: list[str]) -> str:
    for token in tokenize_content_id(content_id):
        if token in sprite_map:
            return sprite_map[token]
    for token in tokenize_content_id(content_id):
        for stem, key in sprite_map.items():
            if token in stem or stem in token:
                return key
    return pick_from_pool(content_id, fallback_pool)


def read_content_visual() -> list[dict]:
    wb = load_workbook(CONTENT_VISUAL_PATH, data_only=True)
    ws = wb[CV_SHEET] if CV_SHEET in wb.sheetnames else wb.active
    rows: list[dict] = []
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = _cell_str(row[0] if row else "")
        if marker.startswith("##"):
            continue
        content_id = _cell_str(row[1] if len(row) > 1 else "")
        if not content_id:
            continue
        rows.append(
            {
                "content_id": content_id,
                "content_kind": _cell_str(row[2] if len(row) > 2 else ""),
                "description": _cell_str(row[3] if len(row) > 3 else ""),
                "face_key": _cell_str(row[4] if len(row) > 4 else ""),
                "frame_key": _cell_str(row[5] if len(row) > 5 else ""),
                "icon_key": _cell_str(row[6] if len(row) > 6 else ""),
                "sheet_row_index": row_index,
            }
        )
    wb.close()
    return rows


def patch_content_visual(patches: list[dict]) -> None:
    wb = load_workbook(CONTENT_VISUAL_PATH)
    ws = wb[CV_SHEET] if CV_SHEET in wb.sheetnames else wb.active
    row_by_id = {}
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = _cell_str(row[0] if row else "")
        if marker.startswith("##"):
            continue
        cid = _cell_str(row[1] if len(row) > 1 else "")
        if cid:
            row_by_id[cid] = row_index

    for patch in patches:
        cid = patch["content_id"]
        row_index = patch.get("sheet_row_index", row_by_id.get(cid, -1))
        if row_index < 0:
            continue
        excel_row = row_index + 1
        if "face_key" in patch:
            ws.cell(row=excel_row, column=5, value=patch.get("face_key") or "")
        if "icon_key" in patch:
            ws.cell(row=excel_row, column=7, value=patch.get("icon_key") or "")

    wb.save(CONTENT_VISUAL_PATH)
    wb.close()


def upsert_visual_assets(upserts: list[dict]) -> None:
    wb = load_workbook(VISUAL_ASSET_PATH)
    ws = wb[VA_SHEET] if VA_SHEET in wb.sheetnames else wb.active
    row_by_id: dict[str, int] = {}
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = _cell_str(row[0] if row else "")
        if marker.startswith("##"):
            continue
        vid = _cell_str(row[1] if len(row) > 1 else "")
        if vid:
            row_by_id[vid] = row_index

    for item in upserts:
        vid = item["visual_id"]
        row_index = item.get("sheet_row_index", row_by_id.get(vid, -1))
        if row_index < 0:
            ws.append(["", vid, item.get("kind", "sprite"), item.get("asset_key", ""), item.get("fallback_id", "")])
            row_by_id[vid] = ws.max_row - 1
            continue
        excel_row = row_index + 1
        ws.cell(row=excel_row, column=3, value=item.get("kind", "sprite"))
        ws.cell(row=excel_row, column=4, value=item.get("asset_key", ""))
        ws.cell(row=excel_row, column=5, value=item.get("fallback_id", ""))

    wb.save(VISUAL_ASSET_PATH)
    wb.close()


def build_assignments(rows: list[dict], force: bool) -> tuple[list[dict], list[dict]]:
    monster_icons = parse_atlas_sprites(PROJECT_ROOT / (MONSTER_ICON_ATLAS + ".meta"))
    if not monster_icons:
        monster_icons = [f"pixel_art_card_pack 1_{i}" for i in range(120)]

    face_keys = [legacy_key(MONSTER_FACE_ATLAS, name) for name in FACE_SPRITE_NAMES]
    monster_icon_keys = [legacy_key(MONSTER_ICON_ATLAS, name) for name in monster_icons]

    item_map = collect_png_sprites(ITEMS_DIR)
    icon_map = collect_png_sprites(ICONS_DIR)
    item_pool = sorted(set(item_map.values()))
    icon_pool = sorted(set(icon_map.values()))

    cv_patches: list[dict] = []
    va_upserts: dict[str, dict] = {}

    def assign(
        row: dict,
        *,
        icon_asset: str = "",
        face_asset: str = "",
    ) -> None:
        cid = row["content_id"]
        patch: dict = {"content_id": cid, "sheet_row_index": row["sheet_row_index"]}
        changed = False

        if icon_asset and (force or not row.get("icon_key")):
            vid = visual_id("icon", cid)
            patch["icon_key"] = vid
            va_upserts[vid] = {"visual_id": vid, "kind": "sprite", "asset_key": icon_asset, "fallback_id": ""}
            changed = True

        if face_asset and (force or not row.get("face_key")):
            vid = visual_id("face", cid)
            patch["face_key"] = vid
            va_upserts[vid] = {"visual_id": vid, "kind": "sprite", "asset_key": face_asset, "fallback_id": ""}
            changed = True

        if changed:
            cv_patches.append(patch)

    for row in rows:
        kind = row["content_kind"]
        cid = row["content_id"]

        if kind == "HelpCard":
            icon = fuzzy_pick(cid, item_map, item_pool)
            face = pick_from_pool(cid, face_keys)
            assign(row, icon_asset=icon, face_asset=face)

        elif kind == "Monster":
            icon = pick_from_pool(cid, monster_icon_keys)
            face = pick_from_pool(cid, face_keys)
            assign(row, icon_asset=icon, face_asset=face)

        elif kind in ("Relic", "Skill"):
            icon = fuzzy_pick(cid, icon_map, icon_pool)
            face = pick_from_pool(cid, face_keys)
            assign(row, icon_asset=icon, face_asset=face)

        elif kind == "MonsterDeck":
            face = pick_from_pool(cid, face_keys)
            assign(row, face_asset=face)

        elif kind == "Room":
            icon = fuzzy_pick(cid, item_map, item_pool)
            assign(row, icon_asset=icon)

        elif kind == "Avatar":
            icon = item_map.get("character") or pick_from_pool(cid, item_pool)
            face = pick_from_pool(cid, face_keys)
            assign(row, icon_asset=icon, face_asset=face)

    return cv_patches, list(va_upserts.values())


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--force", action="store_true", help="Overwrite existing icon/face keys")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if not CONTENT_VISUAL_PATH.exists():
        print(f"Missing {CONTENT_VISUAL_PATH}", file=sys.stderr)
        return 1

    rows = read_content_visual()
    patches, upserts = build_assignments(rows, force=args.force)

    summary = {
        "content_visual_patches": len(patches),
        "visual_asset_upserts": len(upserts),
        "by_kind": {},
    }
    patch_ids = {p["content_id"] for p in patches}
    for row in rows:
        if row["content_id"] in patch_ids:
            summary["by_kind"][row["content_kind"]] = summary["by_kind"].get(row["content_kind"], 0) + 1

    print(json.dumps(summary, ensure_ascii=False, indent=2))

    if args.dry_run:
        return 0

    if patches:
        patch_content_visual(patches)
    if upserts:
        upsert_visual_assets(upserts)

    print("Done. Run gen_table_nine.ps1 to refresh StreamingAssets.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
