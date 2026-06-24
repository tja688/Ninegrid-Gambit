#!/usr/bin/env python3
"""Read / patch card_frame_style.xlsx for Unity ContentVisualEditor."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

SHEET_NAME = "TbCardFrameStyle"
COL_MARKER = 0
COL_STYLE_ID = 1
COL_COLOR_R = 2
COL_COLOR_G = 3
COL_COLOR_B = 4
COL_COLOR_A = 5


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    return float(value)


def read_all(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    rows: list[dict] = []
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = _cell_str(row[COL_MARKER] if len(row) > COL_MARKER else "")
        if marker.startswith("##"):
            continue
        style_id = _cell_str(row[COL_STYLE_ID] if len(row) > COL_STYLE_ID else "")
        if not style_id:
            continue
        rows.append(
            {
                "style_id": style_id,
                "color_r": _cell_float(row[COL_COLOR_R] if len(row) > COL_COLOR_R else None, 1.0),
                "color_g": _cell_float(row[COL_COLOR_G] if len(row) > COL_COLOR_G else None, 1.0),
                "color_b": _cell_float(row[COL_COLOR_B] if len(row) > COL_COLOR_B else None, 1.0),
                "color_a": _cell_float(row[COL_COLOR_A] if len(row) > COL_COLOR_A else None, 1.0),
                "sheet_row_index": row_index,
            }
        )
    wb.close()
    return rows


def patch_colors(path: Path, patches: list[dict]) -> None:
    wb = load_workbook(path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    row_by_id: dict[str, int] = {}
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = _cell_str(row[COL_MARKER] if len(row) > COL_MARKER else "")
        if marker.startswith("##"):
            continue
        style_id = _cell_str(row[COL_STYLE_ID] if len(row) > COL_STYLE_ID else "")
        if style_id:
            row_by_id[style_id] = row_index

    for patch in patches:
        style_id = patch.get("style_id", "")
        if not style_id:
            continue
        sheet_row_index = patch.get("sheet_row_index", -1)
        if sheet_row_index < 0:
            sheet_row_index = row_by_id.get(style_id, -1)
        if sheet_row_index < 0:
            continue
        excel_row = sheet_row_index + 1
        ws.cell(row=excel_row, column=COL_COLOR_R + 1, value=float(patch.get("color_r", 1.0)))
        ws.cell(row=excel_row, column=COL_COLOR_G + 1, value=float(patch.get("color_g", 1.0)))
        ws.cell(row=excel_row, column=COL_COLOR_B + 1, value=float(patch.get("color_b", 1.0)))
        ws.cell(row=excel_row, column=COL_COLOR_A + 1, value=float(patch.get("color_a", 1.0)))

    wb.save(path)
    wb.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["read", "patch"])
    parser.add_argument("--path", required=True)
    parser.add_argument("--patches", default="")
    args = parser.parse_args()
    path = Path(args.path)
    if not path.exists():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    if args.mode == "read":
        print(json.dumps(read_all(path), ensure_ascii=False))
        return 0

    if args.mode == "patch":
        patches = json.loads(args.patches) if args.patches else []
        patch_colors(path, patches)
        return 0

    return 1


if __name__ == "__main__":
    raise SystemExit(main())
