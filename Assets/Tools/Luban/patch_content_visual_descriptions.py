#!/usr/bin/env python3
"""Prepend Chinese display names to TbContentVisual descriptions."""

from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent
DATAS_DIR = SCRIPT_DIR / "Datas"
XLSX_PATH = DATAS_DIR / "content_visual.xlsx"
JSON_PATH = (
    SCRIPT_DIR.parent.parent
    / "StreamingAssets"
    / "TableNine"
    / "LubanData"
    / "tablenine_tbcontentvisual.json"
)

SHEET_NAME = "TbContentVisual"
COL_MARKER = 0
COL_CONTENT_ID = 1
COL_CONTENT_KIND = 2
COL_DESCRIPTION = 3


def load_json(name: str):
    with (DATAS_DIR / name).open(encoding="utf-8") as f:
        return json.load(f)


def build_display_name_map() -> dict[str, str]:
    cards = {row["def_id"]: row.get("display_name", "") for row in load_json("cards.json")}
    relics = {row["def_id"]: row.get("display_name", "") for row in load_json("relics.json")}
    skills = {row["def_id"]: row.get("display_name", "") for row in load_json("skills.json")}
    rooms = {row["kind"]: row.get("display_name", "") for row in load_json("rooms.json")}
    decks = {row["id"]: row.get("display_name", "") for row in load_json("monster_decks.json")}
    names = {}
    names.update(cards)
    names.update(relics)
    names.update(skills)
    names.update(rooms)
    names.update(decks)
    names["avatar.default"] = "玩家"
    return names


def prepend_name(display_name: str, description: str) -> str:
    display_name = (display_name or "").strip()
    description = (description or "").strip()
    if not display_name:
        return description
    if not description:
        return display_name
    if description.startswith(display_name):
        return description
    for sep in ("：", ":", "—", "-", " "):
        prefix = f"{display_name}{sep}"
        if description.startswith(prefix):
            return description
    return f"{display_name}：{description}"


def patch_xlsx(names: dict[str, str]) -> int:
    wb = load_workbook(XLSX_PATH)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    changed = 0
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = str(row[COL_MARKER] or "").strip()
        if marker.startswith("##"):
            continue
        content_id = str(row[COL_CONTENT_ID] or "").strip()
        if not content_id:
            continue
        old_desc = str(row[COL_DESCRIPTION] or "").strip()
        new_desc = prepend_name(names.get(content_id, ""), old_desc)
        if new_desc != old_desc:
            ws.cell(row=row_index + 1, column=COL_DESCRIPTION + 1, value=new_desc)
            changed += 1
    wb.save(XLSX_PATH)
    wb.close()
    return changed


def patch_runtime_json(names: dict[str, str]) -> int:
    with JSON_PATH.open(encoding="utf-8") as f:
        rows = json.load(f)
    changed = 0
    for row in rows:
        content_id = row.get("content_id", "")
        old_desc = (row.get("description") or "").strip()
        new_desc = prepend_name(names.get(content_id, ""), old_desc)
        if new_desc != old_desc:
            row["description"] = new_desc
            changed += 1
    with JSON_PATH.open("w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return changed


def main() -> int:
    names = build_display_name_map()
    xlsx_changed = patch_xlsx(names)
    json_changed = patch_runtime_json(names)
    print(f"Patched xlsx rows: {xlsx_changed}, runtime json rows: {json_changed}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
