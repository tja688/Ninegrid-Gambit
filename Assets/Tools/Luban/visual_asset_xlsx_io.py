#!/usr/bin/env python3
"""Read / upsert visual_asset.xlsx for Unity ContentVisualEditor."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

SHEET_NAME = "TbVisualAsset"
COL_MARKER = 0
COL_VISUAL_ID = 1
COL_KIND = 2
COL_ASSET_KEY = 3
COL_FALLBACK_ID = 4


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_all(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    rows: list[dict] = []
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = _cell_str(row[COL_MARKER] if len(row) > COL_MARKER else "")
        if marker.startswith("##"):
            continue
        visual_id = _cell_str(row[COL_VISUAL_ID] if len(row) > COL_VISUAL_ID else "")
        if not visual_id:
            continue
        rows.append(
            {
                "visual_id": visual_id,
                "kind": _cell_str(row[COL_KIND] if len(row) > COL_KIND else ""),
                "asset_key": _cell_str(row[COL_ASSET_KEY] if len(row) > COL_ASSET_KEY else ""),
                "fallback_id": _cell_str(row[COL_FALLBACK_ID] if len(row) > COL_FALLBACK_ID else ""),
                "sheet_row_index": row_index,
            }
        )
    wb.close()
    return rows


def upsert_rows(path: Path, upserts: list[dict]) -> None:
    wb = load_workbook(path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    row_by_id: dict[str, int] = {}
    for row_index, row in enumerate(ws.iter_rows(values_only=True)):
        marker = _cell_str(row[COL_MARKER] if len(row) > COL_MARKER else "")
        if marker.startswith("##"):
            continue
        visual_id = _cell_str(row[COL_VISUAL_ID] if len(row) > COL_VISUAL_ID else "")
        if visual_id:
            row_by_id[visual_id] = row_index

    for item in upserts:
        visual_id = item.get("visual_id", "")
        if not visual_id:
            continue
        sheet_row_index = item.get("sheet_row_index", -1)
        if sheet_row_index < 0:
            sheet_row_index = row_by_id.get(visual_id, -1)
        if sheet_row_index < 0:
            sheet_row_index = ws.max_row
            ws.append(
                [
                    "",
                    visual_id,
                    item.get("kind", "sprite") or "sprite",
                    item.get("asset_key", "") or "",
                    item.get("fallback_id", "") or "",
                ]
            )
            row_by_id[visual_id] = sheet_row_index
            continue
        excel_row = sheet_row_index + 1
        ws.cell(row=excel_row, column=COL_KIND + 1, value=item.get("kind", "sprite") or "sprite")
        ws.cell(row=excel_row, column=COL_ASSET_KEY + 1, value=item.get("asset_key", "") or "")
        ws.cell(row=excel_row, column=COL_FALLBACK_ID + 1, value=item.get("fallback_id", "") or "")

    wb.save(path)
    wb.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["read", "upsert"])
    parser.add_argument("--path", required=True)
    parser.add_argument("--upserts", default="")
    args = parser.parse_args()
    path = Path(args.path)
    if not path.exists():
        print(f"File not found: {path}", file=sys.stderr)
        return 1

    if args.mode == "read":
        print(json.dumps(read_all(path), ensure_ascii=False))
        return 0

    if args.mode == "upsert":
        upserts = json.loads(args.upserts) if args.upserts else []
        upsert_rows(path, upserts)
        return 0

    return 1


if __name__ == "__main__":
    raise SystemExit(main())
