# -*- coding: utf-8 -*-
"""Probe: xlsx descriptions vs JSON descriptions; idle sprite folder sizes."""
import json
import os
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"c:\Users\jinji\Documents\GitHub\Ninegrid-Gambit")

# ---------- 1. xlsx descriptions ----------
from openpyxl import load_workbook

wb = load_workbook(ROOT / "Assets/Tools/Luban/Datas/content_visual.xlsx", data_only=True)
ws = wb["TbContentVisual"] if "TbContentVisual" in wb.sheetnames else wb.active
xlsx = {}
for r in ws.iter_rows(values_only=True):
    marker = str(r[0] or "")
    if marker.startswith("##"):
        continue
    cid = str(r[1] or "").strip()
    if not cid:
        continue
    xlsx[cid] = str(r[3] or "").strip()

print("=== XLSX rows:", len(xlsx))
for cid in list(xlsx)[:6]:
    print(" ", cid, "|", xlsx[cid][:50])

# ---------- 2. JSON descriptions: detect mojibake ----------
CARDS = ROOT / "Assets/Arts/ContentVisual/cards"

def is_mojibake(s: str) -> bool:
    if not s:
        return False
    bad = sum(1 for ch in s if 0x0100 <= ord(ch) <= 0x07FF or ch == "\ufffd")
    return bad >= max(2, len(s) // 10)

good, bad, empty = [], [], []
jsons = {}
for f in sorted(CARDS.glob("*.json")):
    if f.name == "_index.json":
        continue
    d = json.loads(f.read_text(encoding="utf-8"))
    jsons[d["contentId"]] = d
    desc = d.get("description", "")
    if not desc:
        empty.append(d["contentId"])
    elif is_mojibake(desc):
        bad.append(d["contentId"])
    else:
        good.append((d["contentId"], desc[:30]))

print("\n=== JSON: total", len(jsons), "| mojibake", len(bad), "| good", len(good), "| empty", len(empty))
print("good samples:")
for cid, s in good[:10]:
    print("  ", cid, "|", s)
print("empty:", empty[:20])

# does xlsx have matching (non-mojibake) description for the bad ones?
miss = [cid for cid in bad if cid not in xlsx or not xlsx[cid]]
xbad = [cid for cid in bad if cid in xlsx and is_mojibake(xlsx[cid])]
print("bad-json without xlsx desc:", len(miss), miss[:10])
print("xlsx desc itself mojibake:", len(xbad), xbad[:10])
if bad:
    cid = bad[0]
    print("sample fix:", cid, "| json:", jsons[cid]["description"][:30], "| xlsx:", xlsx.get(cid, "")[:50])

# ---------- 3. deck / kind / anim state overview ----------
from collections import Counter, defaultdict
decks = defaultdict(list)
anim_set = 0
for cid, d in jsons.items():
    decks[(d.get("kind"), d.get("deckId") or "-")].append(cid)
    slots = (d.get("animations") or {}).get("slots") or []
    for s in slots:
        if s.get("id") == "idle" and s.get("sourceType") not in (None, "", "none"):
            anim_set += 1
            print("HAS-IDLE:", cid, s.get("sourceType"), s.get("path"))
print("\n=== decks ===")
for k, v in sorted(decks.items()):
    print(k, len(v))
print("idle already set:", anim_set)

# ---------- 4. sprite folder sizes (idle only) ----------
from PIL import Image

SPR = ROOT / r"Assets/Arts/Images/Png/像素怪物合集/sprites"

def content_bbox(im):
    im = im.convert("RGBA")
    return im.getbbox()  # includes alpha

# baseline
base_dir = SPR / "001insect_attack_01_M"
base_pngs = sorted(base_dir.glob("*.png"))
im = Image.open(base_pngs[0])
bb = content_bbox(im)
print("\n=== baseline 001insect_attack_01_M:", im.size, "content bbox:", bb,
      "content wh:", (bb[2] - bb[0], bb[3] - bb[1]), "frames:", len(base_pngs))

rows = []
for d in sorted(SPR.iterdir()):
    if not d.is_dir():
        continue
    name = d.name
    if "idle" not in name.lower():
        continue
    pngs = sorted(d.glob("*.png"))
    if not pngs:
        continue
    try:
        im = Image.open(pngs[0])
        bb = content_bbox(im)
        cw, ch = (bb[2] - bb[0], bb[3] - bb[1]) if bb else (0, 0)
        rows.append((name, im.size, (cw, ch), len(pngs)))
    except Exception as e:
        rows.append((name, "ERR", str(e), 0))

print("\n=== idle folders:", len(rows))
for name, sz, cwh, n in rows:
    print(f"{name}\t{sz}\t{cwh}\t{n}")
