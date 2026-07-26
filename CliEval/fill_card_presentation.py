# -*- coding: utf-8 -*-
"""一次性数据填充：
1) 61 张怪物卡：idle 动画（像素怪物合集，card_fit usable 集）+ displayName 改为素材中文名
   + uniformScale 用 catalog suggestedUniformScale。
2) 全部卡：乱码 description 从 content_visual.xlsx 回填（怪物卡描述前缀旧名→新名）；
   乱码 displayName 从 Luban（cards/relics/skills.json）回填。
3) authoring(Assets/Arts/ContentVisual/cards) 与 StreamingAssets 双写。
不改 stats/gold/mainIcon/其它槽位；手搓的非乱码文本不覆盖。
"""
import json
import io
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"c:\Users\jinji\Documents\GitHub\Ninegrid-Gambit")
AUTHORING = ROOT / "Assets/Arts/ContentVisual/cards"
STREAMING = ROOT / "Assets/StreamingAssets/ContentVisual/cards"
SPR_ROOT = "Assets/Arts/Images/Png/像素怪物合集/sprites"

# contentId -> (idle folder, 新中文名)
MAPPING = {
    # ---- deck.stone_legion 石头系：矿物/甲壳 ----
    "monster.big_stone":        ("001insect_idle_01", "岩甲虫"),
    "monster.stone_man":        ("21sandmonster_idle_01", "沙岩怪"),
    "monster.megalith":         ("buff_honeycombe_idle_01", "蜂窝岩"),
    "monster.rolling_stone_man": ("38pangolin_idle_01", "黄鳞穿山甲"),
    "monster.sharp_stone":      ("34scorpion_idle_01", "蓝尾蝎"),
    "monster.shelter_stone":    ("54tortoise_idle_01", "棕壳龟"),
    "monster.stone_shrimp":     ("32crab_idle_01", "紫壳蟹"),
    "monster.stone_swallower":  ("55tortoise_idle_01", "红壳龟"),
    "monster.stone_thrower":    ("50ballcactus_idle_01", "球仙人掌"),
    "monster.growing_stone":    ("04mushroom_idle_01", "红伞蘑菇"),
    # ---- deck.skeleton_legion 骷髅系：骷髅/枯木/暗黑虫 ----
    "monster.skull_head":          ("003skull_idle_01", "蓝焰骷髅"),
    "monster.headless_skeleton":   ("004skull_idle_01", "青焰骷髅"),
    "monster.big_skeleton":        ("009skull_idle_01", "棕骨骷髅"),
    "monster.big_skeleton_reborn": ("005skull_idle_01", "蓝纹骷髅"),
    "monster.giant_skeleton":      ("47treemonster_idle", "枯木巨妖"),
    "monster.bone_club_skeleton":  ("30treemonster_idle_01", "枯木妖"),
    "monster.multi_bone_worm":     ("18sandworm_idle_01", "节节沙虫"),
    "monster.bone_courier":        ("10mouse_idle_01", "棕野鼠"),
    "monster.skeleton_taunter":    ("002bat_idle_01", "灰翼蝠"),
    "monster.skeleton_mage":       ("39bigspider_idle_01", "紫背巨蛛"),
    "monster.skeleton_king":       ("Corvyn_idle_01", "紫袍死灵"),
    # ---- deck.orc_legion 兽人系：哥布林/狼骑/野猪 ----
    "monster.young_orc":         ("06goblin_idle_01", "绿皮哥布林"),
    "monster.brainless_orc":     ("08goblin_idle_01", "橙帽哥布林"),
    "monster.smart_orc":         ("07goblin_idle_01", "灰甲哥布林"),
    "monster.orc_warrior":       ("15WolfGobin_idle_01", "狼骑哥布林"),
    "monster.veteran_orc":       ("25WolfGobin_idle_01", "灰狼骑兵"),
    "monster.orc_commander":     ("26wildboardknight_idle_01", "野猪骑士"),
    "monster.old_orc":           ("02wildboar_idle_01", "棕鬃野猪"),
    "monster.big_orc":           ("07bearBoss_attack_idle", "红毛巨熊"),
    "monster.orc_boss":          ("49buffalo_idle_01", "巨角野牛"),
    "monster.orc_quartermaster": ("Buket_idle_01", "红果小贩"),
    # ---- deck.dragon 巨龙系：爬行类/火系 ----
    "monster.fire_dragon":        ("28stegosaurus_idle_01", "赤背剑龙"),
    "monster.salamander":         ("53LizardMan03_idle_01", "棕鳞蜥人"),
    "monster.dragon_follower":    ("51LizardMan01_idle_01", "绿鳞蜥人"),
    "monster.dragon_cult_leader": ("52LizardMan02_idle_01", "绿甲蜥人"),
    "monster.fire_cult_leader":   ("17frog_idle_01", "赤红蛙"),
    "monster.fire_priest":        ("08flower_idle_01", "红焰花"),
    "monster.fire_bather":        ("16walkflower_idle_01", "红瓣行花"),
    "monster.fire_swallower":     ("05mouse_idle_01", "红首鼠"),
    "monster.executioner":        ("03wolf_new_idle_01", "灰鬃狼"),
    "monster.stone_golem":        ("buff_golem_idle_01", "石魔像"),
    # ---- deck.void 虚空系：史莱姆/异形/精灵 ----
    "monster.void_cub":           ("03slime_idle_01", "青史莱姆"),
    "monster.rotating_cub":       ("40slime_idle_01", "绿芽史莱姆"),
    "monster.void_lost":          ("41slime_idle_01", "青纹史莱姆"),
    "monster.mist":               ("Map01_monster01_idle_01", "绿雾团"),
    "monster.observer":           ("09bat_idle_01", "大耳蝠"),
    "monster.sky_eye":            ("13wasp_idle_01", "青翅蜂"),
    "monster.space_master":       ("50flowermonster_idle_01", "幽蓝异花"),
    "monster.world_turning_hand": ("NPC_Dryad_golden_idle", "金树灵"),
    "monster.stepwalker":         ("01rabit_new_idle_02", "竖耳兔"),
    "monster.friendly_ancient":   ("Sum_Elf_golden_idle_01", "金光精灵"),
    # ---- deck.wandering_legion 流浪系：人物/杂兵 ----
    "monster.beggar":          ("Dogge_idle_01", "棕毛土狗"),
    "monster.vagrant":         ("Kyle_idle_01", "绿帽猎人"),
    "monster.wandering_child": ("01rabit_new_idle_01", "棕毛小兔"),
    "monster.pickpocket":      ("Layla_idle_01", "黄帽少女"),
    "monster.killer":          ("02wolf_idle_01", "灰狼"),
    "monster.rogue":           ("Philip_idle_01", "蓝帽枪手"),
    "monster.thug":            ("Icey_idle_01", "白发刀客"),
    "monster.smuggler":        ("NPC_smithy_idle", "围裙铁匠"),
    "monster.hoodlum":         ("NPC_tavern_idle01", "酒馆汉子"),
    "monster.ringleader":      ("Musashi_idle_01", "蓝甲武士"),
}


def is_mojibake(s: str) -> bool:
    if not s:
        return False
    bad = sum(1 for ch in s if 0x0100 <= ord(ch) <= 0x07FF or ch == "\ufffd")
    return bad >= max(2, len(s) // 10)


def load_xlsx_desc():
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "Assets/Tools/Luban/Datas/content_visual.xlsx", data_only=True)
    ws = wb["TbContentVisual"] if "TbContentVisual" in wb.sheetnames else wb.active
    m = {}
    for r in ws.iter_rows(values_only=True):
        if str(r[0] or "").startswith("##"):
            continue
        cid = str(r[1] or "").strip()
        if cid:
            m[cid] = str(r[3] or "").strip()
    return m


def load_luban_names():
    m = {}
    datas = ROOT / "Assets/Tools/Luban/Datas"
    for fn, key in [("cards.json", "def_id"), ("relics.json", "def_id"), ("skills.json", "def_id")]:
        for rec in json.loads((datas / fn).read_text(encoding="utf-8")):
            m[rec[key]] = rec.get("display_name", "")
    return m


def load_catalog_scales():
    cat = json.loads(
        (ROOT / SPR_ROOT / "_card_fit/catalog.json").read_text(encoding="utf-8"))
    m = {}
    for c in cat["clips"]:
        m[c["folder"]] = (bool(c.get("usable")), round(float(c.get("suggestedUniformScale", 1.0)), 2))
    return m


def main():
    dry = "--dry" in sys.argv
    xlsx = load_xlsx_desc()
    luban = load_luban_names()
    scales = load_catalog_scales()

    # 校验映射素材全部存在且 usable
    errs = []
    for cid, (folder, _name) in MAPPING.items():
        if not (ROOT / SPR_ROOT / folder).is_dir():
            errs.append(f"missing folder: {cid} -> {folder}")
        elif folder not in scales:
            errs.append(f"not in catalog: {cid} -> {folder}")
        elif not scales[folder][0]:
            errs.append(f"not usable: {cid} -> {folder}")
    if errs:
        print("\n".join(errs))
        sys.exit(1)

    stats = {"anim": 0, "rename": 0, "desc_fixed": 0, "name_fixed": 0, "skipped_desc": 0}
    for f in sorted(AUTHORING.glob("*.json")):
        if f.name == "_index.json":
            continue
        raw = f.read_text(encoding="utf-8")
        d = json.loads(raw)
        cid = d["contentId"]
        changed = False
        old_luban_name = luban.get(cid, "")

        # --- 名字 ---
        if cid in MAPPING:
            new_name = MAPPING[cid][1]
            if d.get("displayName") != new_name:
                d["displayName"] = new_name
                stats["rename"] += 1
                changed = True
        elif old_luban_name and (
            is_mojibake(d.get("displayName", "")) or d.get("displayName", "") == cid
        ) and d.get("displayName") != old_luban_name:
            d["displayName"] = old_luban_name
            stats["name_fixed"] += 1
            changed = True

        # --- 描述 ---
        if is_mojibake(d.get("description", "")):
            desc = xlsx.get(cid, "")
            if desc:
                if cid in MAPPING:
                    new_name = MAPPING[cid][1]
                    if old_luban_name and desc.startswith(old_luban_name + "："):
                        desc = new_name + "：" + desc[len(old_luban_name) + 1:]
                    elif old_luban_name and desc == old_luban_name:
                        desc = new_name
                d["description"] = desc
                stats["desc_fixed"] += 1
                changed = True
            else:
                stats["skipped_desc"] += 1
                print(f"WARN no xlsx desc for {cid}")

        # --- idle 动画 + 缩放 ---
        if cid in MAPPING:
            folder, _ = MAPPING[cid]
            path = f"{SPR_ROOT}/{folder}"
            slots = d.setdefault("animations", {}).setdefault("slots", [])
            idle = next((s for s in slots if s.get("id") == "idle"), None)
            if idle is None:
                idle = {"id": "idle", "sourceType": "none", "path": "", "offsetX": 0.0, "offsetY": 0.0}
                slots.insert(0, idle)
            if idle.get("sourceType") != "folder" or idle.get("path") != path:
                idle["sourceType"] = "folder"
                idle["path"] = path
                idle["offsetX"] = 0.0
                idle["offsetY"] = 0.0
                stats["anim"] += 1
                changed = True
            mv = d.setdefault("mainVisual", {"offsetX": 0.0, "offsetY": 0.0, "uniformScale": 1.0})
            scale = scales[folder][1]
            if abs(float(mv.get("uniformScale", 1.0)) - scale) > 1e-4:
                mv["uniformScale"] = scale
                changed = True

        if changed and not dry:
            newline = "\r\n" if "\r\n" in raw else "\n"
            text = json.dumps(d, ensure_ascii=False, indent=4)
            if newline == "\r\n":
                text = text.replace("\n", "\r\n")
            data = (text + newline).encode("utf-8")
            f.write_bytes(data)
            sf = STREAMING / f.name
            sf.parent.mkdir(parents=True, exist_ok=True)
            sf.write_bytes(data)

    print("dry-run" if dry else "written", stats)


if __name__ == "__main__":
    main()
